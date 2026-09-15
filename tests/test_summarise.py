import pandas as pd
import pytest
from docx import Document as open_document
from docx.document import Document
from great_tables import GT

import pyfixest as pf
from pyfixest.estimation import feols, fepois
from pyfixest.report.summarize import etable, summary
from pyfixest.utils.dev_utils import _select_order_coefs
from pyfixest.utils.utils import get_data


def test_summary():
    """Just run etable() and summary() on a few models."""
    df1 = get_data()
    df1 = pd.concat(
        [df1, df1], axis=0
    )  # Make it a bit larger, for examining the thousands separator
    df2 = get_data(model="Fepois")

    fit1 = feols("Y ~ X1 + X2 | f1", data=df1)
    fit1a = feols("Y ~ X1 + X2 + f1", data=df1)
    fit2 = fepois("Y ~ X1 + X2 + f2 | f1", data=df2, vcov={"CRV1": "f1+f2"})
    fit3 = feols("Y ~ X1", data=df1)
    fit4 = feols("Y ~ X1", data=df1, weights="weights")
    fit5 = feols("Y ~ 1 | Z1 ~ X1", data=df1)

    fit_qreg = pf.quantreg("Y ~ X1", data=df1, vcov="nid")

    summary(fit1)
    summary(fit2)
    summary([fit1, fit2])
    summary([fit4])
    fit5.summary()

    etable(fit1)
    etable(fit2)
    etable([fit1, fit2])

    etable([fit3])
    etable([fit1, fit2, fit3])

    fit_iv = feols("Y ~ X2 | f1 | X1 ~ Z1", data=df1)
    etable([fit_iv, fit1])

    fit_multi = feols("Y + Y2 ~ X1 + X2 | f1", data=df1)
    etable(fit_multi.to_list())

    # Test significance code
    etable([fit1, fit2], signif_code=[0.01, 0.05, 0.1])
    etable([fit1, fit2], signif_code=[0.02, 0.06, 0.1])

    # Test coefficient format
    etable([fit1, fit2], coef_fmt="b (se)\nt [p]")

    # Test custom statistics
    etable(
        models=[fit1, fit2],
        custom_stats={
            "conf_int_lb": [fit1._conf_int[0], fit2._conf_int[0]],
            "conf_int_ub": [fit1._conf_int[1], fit2._conf_int[1]],
        },
        coef_fmt="b [conf_int_lb, conf_int_ub]",
    )

    # Test scientific notation
    etable(
        models=[fit1],
        custom_stats={
            "test_digits": [[0.1, 12300]],
        },
        coef_fmt="b [test_digits]",
        digits=2,
    )

    # Test scientific notation, thousands separator
    etable(
        models=[fit1],
        custom_stats={
            "test_digits": [[0.1, 12300]],
        },
        coef_fmt="b [test_digits]",
        digits=2,
        scientific_notation=False,
        thousands_sep=True,
    )

    # Test select / order coefficients
    etable([fit1, fit2, fit3], coef_fmt="b (se)\nt [p]")
    etable([fit1, fit2, fit3], coef_fmt="b (se)\nt [p]", keep=["X1", "cep"])
    etable([fit1, fit2, fit3], coef_fmt="b (se)\nt [p]", drop=[r"\d$"])
    etable([fit1, fit2, fit3], coef_fmt="b (se)\nt [p]", keep=[r"\d"], drop=["f"])
    etable([fit1, fit2, fit3], coef_fmt="b (se)\nt [p]", keep="X")
    etable([fit1, fit2, fit3], coef_fmt="b (se)\nt [p]", drop=r"\d$")

    # test labels, felabels args
    etable([fit1, fit1a], labels={"X1": "X1_label"}, felabels={"f1": "f1_label"})
    etable(
        [fit1, fit1a], labels={"X1": "X1_label"}, felabels={"f1": "f1_label"}, keep="X1"
    )
    etable(
        [fit1, fit1a], labels={"X1": "X1_label"}, felabels={"f1": "f1_label"}, drop="X1"
    )
    etable([fit1, fit1a], felabels={"f1": "f1_renamed2"}, keep=["f1"])

    cols = ["x1", "x2", "x11", "x21"]
    assert _select_order_coefs(cols, keep=["x1"]) == ["x1", "x11"]
    assert _select_order_coefs(cols, drop=["x1"]) == ["x2", "x21"]
    assert _select_order_coefs(cols, keep=["x1"], exact_match=True) == ["x1"]
    assert _select_order_coefs(cols, drop=["x1"], exact_match=True) == [
        "x2",
        "x11",
        "x21",
    ]

    # API tests for new tex args

    etable([fit1, fit2], type="tex")

    etable([fit1, fit2], type="tex", notes="You can add notes here.")
    etable([fit1, fit2], type="md", notes="You can add notes here.")

    etable([fit1, fit2], type="tex", model_heads=["Model 1", "Model 2"])
    etable(
        [fit1, fit2], type="tex", model_heads=["Model 1", "Model 2"], head_order="dh"
    )
    etable(
        [fit1, fit2], type="tex", model_heads=["Model 1", "Model 2"], head_order="hd"
    )
    etable([fit1, fit2], type="tex", model_heads=["Model 1", "Model 2"], head_order="d")
    etable([fit1, fit2], type="tex", model_heads=["Model 1", "Model 2"], head_order="h")
    etable([fit1, fit2], type="tex", model_heads=["Model 1", "Model 2"], head_order="")
    etable([fit1, fit2], type="tex", file_name="tests/texfiles/test.tex")

    summary(fit_qreg)
    etable(fit_qreg)


def test_etable_significance_stars_follow_coef_fmt():
    data = get_data()
    fit1 = feols("Y ~ X1", data=data)
    fit2 = feols("Y ~ X1 + X2 | f1", data=data)

    default_table = etable([fit1, fit2], type="df")
    assert any("*" in str(value) for value in default_table.to_numpy().ravel())

    custom_fmt_table = etable([fit1, fit2], type="df", coef_fmt="b (se)\nt [p]")
    assert not any("*" in str(value) for value in custom_fmt_table.to_numpy().ravel())


def test_etable_correct_output_type():
    data = get_data()
    fit = feols("Y ~ X1", data=data)

    df_table = pf.etable(fit, type="df")
    assert isinstance(df_table, pd.DataFrame)

    md_table = pf.etable(fit, type="md")
    assert md_table is None

    gt_table = pf.etable(fit, type="gt")
    assert isinstance(gt_table, GT)

    tex_table = pf.etable(fit, type="tex")
    assert isinstance(tex_table, str)

    typst_table = pf.etable(fit, type="typst")
    assert isinstance(typst_table, str)

    docx_table = pf.etable(fit, type="docx")
    assert isinstance(docx_table, Document)


@pytest.mark.parametrize(
    "formula, fit_kwargs",
    [
        ("Y ~ X1 + X2 | f1", {}),
        ("sw(Y, Y2) ~ X1 + X2 | f1", {}),
        ("Y ~ X2 + [X1 ~ Z1] | f1", {}),
        ("Y ~ X1 + X2 | f1", {"weights": "weights"}),
        ("Y ~ X1 + X2 | f1", {"weights": "frequency", "weights_type": "fweights"}),
        ("Y ~ X1 + X2 | f1", {"lean": True}),
        ("Y ~ X1 + X2 | f1", {"store_data": False}),
    ],
    ids=["fe", "multi", "iv", "aweights", "fweights", "lean", "no-data"],
)
def test_etable_docx_roundtrip(formula, fit_kwargs, tmp_path):
    """Word preserves etable's displayed cells for different fitted-model inputs."""
    data = get_data(N=150, seed=42)
    data["frequency"] = [1 + i % 3 for i in range(len(data))]
    fit = feols(formula, data=data, **fit_kwargs)
    options = dict(
        coef_fmt="b:.3f*\n(se:.3f)",
        signif_code=[0.01, 0.05, 0.1],
        labels={"X1": "政策暴露", "X2": "Control"},
        notes="注: 模拟数据。",
    )
    expected = etable(fit, type="df", **options)
    path = tmp_path / "regression.docx"
    document = etable(fit, type="docx", file_name=path, **options)
    reopened = open_document(path)
    assert isinstance(document, Document)
    assert reopened._element.xml == document._element.xml
    assert len(reopened.tables) == 1
    assert len(reopened.inline_shapes) == 0
    rows = {row.cells[0].text: row for row in reopened.tables[0].rows}
    for label in ("政策暴露", "Control"):
        values = expected.xs(label, level=-1).iloc[0].tolist()
        assert [cell.text for cell in rows[label].cells[1:]] == values
    assert reopened.tables[0].rows[-1].cells[0].text == options["notes"]


def test_etable_docx_style_and_existing_outputs():
    """Word styles are per-call and leave the model and other output formats unchanged."""
    data = get_data(N=150, seed=42)
    models = [feols("Y ~ X1 | f1", data=data), feols("Y ~ X1 + X2 | f1", data=data)]
    options = dict(model_heads=["Baseline", "Controls"], notes="Clustered by firm.")
    before_df = etable(models, type="df", **options)
    before_tex = etable(models, type="tex", **options)
    style = {"font_name": "Arial", "font_size_pt": 12, "notes_font_size_pt": 8}
    document = etable(models, type="docx", docx_style=style, **options)
    table = document.tables[0]
    assert table.rows[0].cells[1]._tc is table.rows[0].cells[2]._tc
    assert [cell.text for cell in table.rows[1].cells[1:]] == ["Baseline", "Controls"]
    assert [cell.text for cell in table.rows[2].cells[1:]] == ["(1)", "(2)"]
    run = table.rows[1].cells[1].paragraphs[0].runs[0]
    assert run.font.name == "Arial"
    assert run.font.size.pt == 12
    notes_run = table.rows[-1].cells[0].paragraphs[0].runs[0]
    assert notes_run.font.size.pt == 8
    assert style == {"font_name": "Arial", "font_size_pt": 12, "notes_font_size_pt": 8}
    pd.testing.assert_frame_equal(etable(models, type="df", **options), before_df)
    assert etable(models, type="tex", **options) == before_tex


def test_etable_invalid_output_type():
    """The invalid-format message includes both Word and Typst output."""
    with pytest.raises(AssertionError, match="'typst' or 'docx'"):
        etable([], type="invalid")


def test_dtable_is_not_public():
    assert "dtable" not in pf.__all__
    assert not hasattr(pf, "dtable")
    assert "dtable" not in pf.report.__all__
    assert not hasattr(pf.report, "dtable")


def test_summary_inference_type_regular(capsys):
    """summary(inference_type='regular') matches the default summary output."""
    fit = feols("Y ~ X1 + X2 | f1", data=get_data())

    summary(fit)
    default_out = capsys.readouterr().out
    summary(fit, inference_type="regular")
    regular_out = capsys.readouterr().out

    assert regular_out
    assert regular_out == default_out
